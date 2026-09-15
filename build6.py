"""The Voter Turnout sheet, from Kalshi's threshold ladder.

Kalshi states turnout as THRESHOLD markets ("Above 340K"), so each rung is
P(value >= strike), a survival curve rather than a set of mutually exclusive
brackets. Differencing adjacent rungs turns it back into buckets.

Kalshi also quotes a margin-of-victory ladder (KXMIDTERMMOV-PA07D / -PA07R).
collect3.py still collects it into data3.json every run, so the history is
kept, but it is deliberately not built into the workbook or published on the
site: as of mid-September 2026 four of the five Democratic rungs had no
two-sided quote, and an expected margin summed over the one rung that did
read as a confident number computed from a fifth of a distribution. The
section that built it is in git history (before 15 Sep 2026) if the ladder
ever fills in.
"""
import json, datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

D3 = json.load(open("data3.json"))
ASOF = datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")

F = "Arial"
BLUE = Font(name=F, size=10, color="0000FF"); BLACK = Font(name=F, size=10)
GREEN = Font(name=F, size=10, color="008000"); BOLD = Font(name=F, size=10, bold=True)
TITLE = Font(name=F, size=16, bold=True, color="1F3864")
SUB = Font(name=F, size=9, italic=True, color="595959")
SECT = Font(name=F, size=11, bold=True, color="FFFFFF")
HDR = Font(name=F, size=9, bold=True, color="FFFFFF")
NOTE = Font(name=F, size=9, italic=True, color="595959")
SECTF = PatternFill("solid", fgColor="1F3864"); HDRF = PatternFill("solid", fgColor="4472C4")
DEMF = PatternFill("solid", fgColor="DEEBF7"); REPF = PatternFill("solid", fgColor="FCE4E4")
YEL = PatternFill("solid", fgColor="FFFF00"); GREY = PatternFill("solid", fgColor="F2F2F2")
BAND = PatternFill("solid", fgColor="F7F9FC")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
PCT = '0.0%'; PCT2 = '0.00%'; DIFF = '+0.0%;-0.0%;0.0%'; DATE = 'yyyy-mm-dd'
NUM = '#,##0;(#,##0);-'; PTS = '+0.00" pts";-0.00" pts";0.00" pts"'
KVOT = '#,##0,"K"'

# Open-ended rungs have no upper edge, so their representative value is a
# judgement call. Everything else uses the true centre of the bucket.
TURN_LOW_MID = 295000   # below the lowest quoted threshold (310K)
TURN_TOP_MID = 385000   # above the highest quoted threshold (370K)
TURNOUT_2024 = 403314   # actual 2024 PA-07 House vote, Mackenzie 50.5 / Wild 49.5

wb = load_workbook("_stage2b.xlsx")


def section(ws, row, text, width):
    ws.cell(row, 1, text).font = SECT
    for c in range(1, width + 1): ws.cell(row, c).fill = SECTF
    return row + 1


def headers(ws, row, cols):
    for i, h in enumerate(cols, 1):
        c = ws.cell(row, i, h); c.font = HDR; c.fill = HDRF; c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30
    return row + 1


def put(ws, r, c, v, font=BLACK, fmt=None, fill=None, align=None, border=True):
    cell = ws.cell(r, c, v); cell.font = font
    if fmt: cell.number_format = fmt
    if fill: cell.fill = fill
    if align: cell.alignment = Alignment(horizontal=align)
    if border: cell.border = BOX
    return cell


def note(ws, r, text, span):
    c = ws.cell(r, 1, text); c.font = NOTE
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
    return r + 1


def need(n, what):
    if n < 1:
        raise SystemExit("FATAL: %s came back empty — refusing to write an inverted row "
                         "range. Re-run collect3.py before building." % what)
    return n


# =====================================================================
# SHEET: Voter Turnout
# =====================================================================
tn = wb.create_sheet("Voter Turnout")
tn.sheet_view.showGridLines = False
tn["A1"] = "Voter Turnout — PA-07 (Kalshi)"; tn["A1"].font = TITLE
tn["A2"] = ("Kalshi KXMIDTERMVOTETURN-PA07: how many ballots the PA-07 House race draws on "
            "3 Nov 2026, quoted as thresholds. Pulled %s." % ASOF)
tn["A2"].font = SUB
for k, v in {"A": 20, "B": 15, "C": 12, "D": 12, "E": 17, "F": 18, "G": 14,
             "H": 18, "I": 15, "J": 13}.items():
    tn.column_dimensions[k].width = v

r = 4
r = section(tn, r, "MARKET METADATA", 10)
T = D3["turnout"]
for lbl, val in [("Event", T["event"]), ("Series", T.get("series")),
                 ("Question", T.get("title")), ("Resolution", "3 Nov 2026"),
                 ("Thresholds quoted", len(T["rungs"]))]:
    put(tn, r, 1, lbl, BOLD); put(tn, r, 2, val, BLUE, NUM if lbl == "Thresholds quoted" else None)
    tn.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7); r += 1
r += 1

r = section(tn, r, "THRESHOLD LADDER", 10)
TH = headers(tn, r, ["Rung", "Threshold", "YES Bid", "YES Ask", "Mid P(turnout ≥ strike)",
                     "Bucket", "Bucket Prob", "Representative Turnout", "Prob × Turnout", ""])
T0 = TH
rungs = sorted(T["rungs"], key=lambda x: x["strike"])
need(len(rungs), "turnout ladder")
# Below the lowest quoted threshold is not a market, it is the residual.
rr = TH
put(tn, rr, 1, "Below %s" % ("{:,.0f}K".format(rungs[0]["strike"] / 1000)), BOLD, None, GREY)
put(tn, rr, 6, "under %.0fK" % (rungs[0]["strike"] / 1000), BLACK)
put(tn, rr, 7, '=IFERROR(MAX(0,1-E{0}),"")'.format(TH + 1), BLACK, PCT2)
put(tn, rr, 8, TURN_LOW_MID, BLUE, KVOT)
put(tn, rr, 9, '=IFERROR(G{0}*H{0},"")'.format(rr), BLACK, '#,##0')
tn.cell(rr, 7).comment = Comment(
    "Not a quoted market — the residual, 1 minus P(turnout above the lowest threshold).", "Tracker")
for i, rung in enumerate(rungs):
    rr = TH + 1 + i
    nxt = rungs[i + 1] if i + 1 < len(rungs) else None
    put(tn, rr, 1, rung["label"], BOLD)
    put(tn, rr, 2, rung["strike"], BLUE, KVOT)
    put(tn, rr, 3, rung.get("yes_bid"), BLUE, PCT)
    put(tn, rr, 4, rung.get("yes_ask"), BLUE, PCT)
    put(tn, rr, 5, '=IF(COUNT(C{0}:D{0})=2,(C{0}+D{0})/2,"")'.format(rr), BLACK, PCT2)
    if nxt:
        put(tn, rr, 6, "%.0fK–%.0fK" % (rung["strike"] / 1000, nxt["strike"] / 1000), BLACK)
        put(tn, rr, 7, '=IFERROR(MAX(0,E{0}-E{1}),"")'.format(rr, rr + 1), BLACK, PCT2)
        put(tn, rr, 8, (rung["strike"] + nxt["strike"]) / 2.0, BLUE, KVOT)
    else:
        put(tn, rr, 6, "above %.0fK" % (rung["strike"] / 1000), BLACK)
        put(tn, rr, 7, '=IFERROR(E{0},"")'.format(rr), BLACK, PCT2)
        put(tn, rr, 8, TURN_TOP_MID, BLUE, KVOT)
    put(tn, rr, 9, '=IFERROR(G{0}*H{0},"")'.format(rr), BLACK, '#,##0')
TN = TH + len(rungs)
TTOT = TN + 1
put(tn, TTOT, 1, "TOTAL", BOLD, None, GREY)
put(tn, TTOT, 7, '=SUM(G{0}:G{1})'.format(T0, TN), BOLD, PCT2, YEL)
put(tn, TTOT, 9, '=SUM(I{0}:I{1})'.format(T0, TN), BOLD, '#,##0')
tn.cell(TH - 1, 8).comment = Comment(
    "ASSUMPTION: the representative turnout for each bucket. Closed buckets use their true "
    "centre. The two open-ended ones are judgement calls — below the lowest threshold is taken "
    "as %s and above the highest as %s. Both feed the expected-turnout figure; change them here "
    "and it updates." % ("{:,}".format(TURN_LOW_MID), "{:,}".format(TURN_TOP_MID)), "Tracker")
r = TTOT + 2

r = section(tn, r, "IMPLIED READ", 10)
put(tn, r, 1, "Expected turnout", BOLD)
put(tn, r, 2, '=IFERROR(SUM(I{0}:I{1}),"")'.format(T0, TN), BOLD, '#,##0', YEL); TEXP = r; r += 1
put(tn, r, 1, "Modal bucket", BOLD)
put(tn, r, 2, '=IFERROR(INDEX($F${0}:$F${1},MATCH(MAX($G${0}:$G${1}),$G${0}:$G${1},0)),"")'.format(T0, TN), BOLD); r += 1
put(tn, r, 1, "P(turnout above 340K)", BOLD)
put(tn, r, 2, '=IFERROR(INDEX($E${0}:$E${1},MATCH(340000,$B${0}:$B${1},0)),"")'.format(T0, TN), BLACK, PCT2); r += 1
put(tn, r, 1, "2024 PA-07 House turnout", BOLD); put(tn, r, 2, TURNOUT_2024, BLUE, '#,##0'); T24 = r; r += 1
put(tn, r, 1, "Implied change vs 2024", BOLD)
put(tn, r, 2, '=IFERROR(B{0}/B{1}-1,"")'.format(TEXP, T24), BOLD, DIFF); r += 1
tn.cell(T24, 2).comment = Comment(
    "403,314 ballots were cast in the 2024 PA-07 House race (Mackenzie 50.5%% / Wild 49.5%%), "
    "per the certified count as reported by Ballotpedia and AP. NOTE: 2024 was a presidential "
    "year and 2026 is a midterm, so a double-digit fall here is the normal pattern, not a "
    "collapse — do not read the change line as a turnout warning.", "Tracker")
r = note(tn, r, "Turnout is a genuine driver here, not trivia: PA-07 was decided by about a point in "
                "2024, so which electorate shows up plausibly decides the seat.", 10)
r += 1

r = section(tn, r, "DAILY HISTORY (mid of closing bid and ask)", 10)
th = D3.get("turnout_history") or {}
tks = [x["ticker"] for x in rungs]
tlabels = [x["label"] for x in rungs]
TH2 = headers(tn, r, ["Date"] + tlabels + [""] * max(0, 9 - len(tlabels)))
tdays = sorted(th)
for i, day in enumerate(tdays):
    rr = TH2 + i
    put(tn, rr, 1, datetime.datetime.strptime(day, "%Y-%m-%d"), BLUE, DATE)
    for j, tk in enumerate(tks):
        rec = (th.get(day) or {}).get(tk)
        v = None
        if rec:
            b, a = rec.get("bid"), rec.get("ask")
            if b is not None and a is not None: v = (b + a) / 2
            elif a is not None: v = a
            elif b is not None: v = b
        put(tn, rr, 2 + j, v, BLUE, PCT)
    if i % 2:
        for c in range(1, len(tks) + 2): tn.cell(rr, c).fill = BAND
TF = TH2
TL = TH2 + need(len(tdays), "turnout history") - 1
r = TL + 1
r = note(tn, r, "Columns are nested, not exclusive: each is P(turnout ≥ that threshold), so within a "
                "row they must fall from left to right. Kalshi began quoting this ladder on %s."
                % (min(tdays) if tdays else "n/a"), 10)
tn.freeze_panes = tn.cell(TH2, 2)

# The turnout midpoints travel out with the row anchors: export_site.py has to
# reproduce this ladder's expected turnout, and a second copy of the constants
# in a second file is a guaranteed future disagreement between the sheet and
# the site.
json.dump({"TURN_LOW_MID": TURN_LOW_MID, "TURN_TOP_MID": TURN_TOP_MID,
           "TURNOUT_2024": TURNOUT_2024,
           "T0": T0, "TN": TN, "TTOT": TTOT, "TEXP": TEXP,
           "TF": TF, "TL": TL, "N_TURN": len(rungs)},
          open("k3refs.json", "w"))
wb.save("_stage2c.xlsx")
print("stage2c ok — turnout rows %d-%d, hist %d days" % (T0, TN, len(tdays)))
