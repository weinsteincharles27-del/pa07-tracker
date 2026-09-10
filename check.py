"""Structural audit of every formula and every chart in the workbook.

Verifies, without a spreadsheet engine:
  1. every cross-sheet reference names a sheet that exists
  2. every referenced cell/range lies inside the sheet's used area, and no
     reference addresses a row above row 1
  3. parentheses and quotes balance
  4. no function outside the LibreOffice/Excel-2007 safe set
  5. self-referencing cells (accidental circularity)
  6. every single-cell reference to an empty cell sits inside a blank guard
  7. every chart series: one non-empty header cell as its title, value and
     category ranges of equal length, no inverted range, no missing sheet

Check 6 is the one that matters most. A reference to an empty cell is not an
error in Excel — it reads as 0 — so it never shows up as #VALUE! or #REF!. It
just quietly becomes a price of 0.000 or a $0 volume. The build scripts guard
against that with COUNT / ISNUMBER / ISBLANK or an explicit ="" or =0 test, and
this check insists that every empty target actually has one. IFERROR does not
count: it catches errors, not empties, which is exactly how INDEX onto a blank
cell shipped as a zero.
"""
import re, sys, json
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

WB = "PA-07_House_Election_Tracker.xlsx"
wb = load_workbook(WB)
sheets = {s.title: s for s in wb.worksheets}

SAFE = {
    "SUM","AVERAGE","MIN","MAX","COUNT","COUNTA","IF","IFERROR","INDEX","MATCH","ABS",
    "ROW","COLUMN","TEXT","DATE","TODAY","AND","OR","NOT","ISNUMBER","ISBLANK","SUMIF",
    "SUMIFS","COUNTIF","COUNTIFS","SUMPRODUCT","ROUND","MEDIAN","STDEV","LEN","CONCATENATE",
}
BANNED = {"XLOOKUP","XMATCH","SORT","FILTER","UNIQUE","SEQUENCE","TEXTJOIN","LET","LAMBDA"}

CELL = re.compile(
    r"(?:(?:'(?P<qs>[^']+)'|(?P<us>[A-Za-z_][A-Za-z0-9_ ]*))!)?"
    r"\$?(?P<c1>[A-Z]{1,3})\$?(?P<r1>\d+)"
    r"(?::\$?(?P<c2>[A-Z]{1,3})\$?(?P<r2>\d+))?"
)
FUNC = re.compile(r"([A-Z][A-Z0-9_\.]*)\s*\(")
# a column letter glued straight onto a minus sign, e.g. the B-5 that row-offset
# arithmetic used to emit off the top of a data block. CELL cannot see these:
# they are not legal references, so they never match, and they never get checked.
MALFORMED = re.compile(r"(?<![A-Z0-9_$])[A-Z]{1,3}-\d")
# the four idioms the build scripts use to turn a blank into a blank, not a zero
GUARDFN = re.compile(r"\b(?:COUNT|COUNTA|ISNUMBER|ISBLANK)\s*\(")

dims = {}
for name, ws in sheets.items():
    dims[name] = (ws.max_row, ws.max_column)

problems = defaultdict(list)
stats = defaultdict(int)


def span(m, home):
    """(sheet, r1, c1, r2, c2) for a CELL match, defaulting to the home sheet."""
    tgt = m.group("qs") or m.group("us") or home
    c1 = column_index_from_string(m.group("c1")); r1 = int(m.group("r1"))
    c2 = column_index_from_string(m.group("c2")) if m.group("c2") else c1
    r2 = int(m.group("r2")) if m.group("r2") else r1
    return tgt, r1, c1, r2, c2


def args_of(f, lparen):
    """Text between the parentheses of a call whose '(' sits at index lparen."""
    depth = 0
    for i in range(lparen, len(f)):
        if f[i] == "(":
            depth += 1
        elif f[i] == ")":
            depth -= 1
            if depth == 0:
                return f[lparen + 1:i]
    return ""


def guarded(f, home):
    """Every cell this formula explicitly tests for emptiness.

    COUNT / COUNTA / ISNUMBER / ISBLANK over a reference, and an explicit ="" or
    =0 comparison against one, are the ways the build scripts declare "this may
    be blank, and blank is handled". Anything a formula reads without one of
    these is being trusted to hold a number.
    """
    out = set()
    for g in GUARDFN.finditer(f):
        for m in CELL.finditer(args_of(f, g.end() - 1)):
            tgt, r1, c1, r2, c2 = span(m, home)
            for rr in range(r1, r2 + 1):
                for cc in range(c1, c2 + 1):
                    out.add((tgt, rr, cc))
    for m in CELL.finditer(f):
        tail = f[m.end():m.end() + 3]
        if tail.startswith('=""') or tail.startswith("=0"):
            tgt, r1, c1, r2, c2 = span(m, home)
            for rr in range(r1, r2 + 1):
                for cc in range(c1, c2 + 1):
                    out.add((tgt, rr, cc))
    return out


def is_empty(tgt, r, c):
    v = sheets[tgt].cell(r, c).value
    return v is None or (isinstance(v, str) and not v.strip())


# ---------------------------------------------------------------- formulas
for name, ws in sheets.items():
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if not (isinstance(v, str) and v.startswith("=")):
                continue
            stats["formulas"] += 1
            f = v[1:]

            if f.count("(") != f.count(")"):
                problems["unbalanced_parens"].append(f"{name}!{cell.coordinate}: {v[:70]}")
            if f.count('"') % 2:
                problems["unbalanced_quotes"].append(f"{name}!{cell.coordinate}: {v[:70]}")

            # functions
            for fn in FUNC.findall(f):
                base = fn.replace("_xlfn.", "")
                if base in BANNED:
                    problems["banned_function"].append(f"{name}!{cell.coordinate}: {fn}")
                elif base not in SAFE:
                    problems["unknown_function"].append(f"{name}!{cell.coordinate}: {fn}")

            # strip string literals before scanning refs
            bare = re.sub(r'"[^"]*"', '""', f)
            for m in MALFORMED.finditer(bare):
                problems["malformed_ref"].append(f"{name}!{cell.coordinate}: {m.group(0)} in {v[:60]}")

            guards = None
            for m in CELL.finditer(bare):
                tgt, r1, c1, r2, c2 = span(m, name)
                if tgt not in sheets:
                    if tgt.upper() not in ("TRUE", "FALSE"):
                        problems["bad_sheet_ref"].append(f"{name}!{cell.coordinate} -> '{tgt}'")
                    continue
                mr, mc = dims[tgt]
                if r1 < 1 or c1 < 1:
                    problems["ref_above_sheet"].append(
                        f"{name}!{cell.coordinate} -> {tgt}!{m.group('c1')}{m.group('r1')}")
                    continue
                if r2 > mr or c2 > mc:
                    problems["ref_out_of_range"].append(
                        f"{name}!{cell.coordinate} -> {tgt}!{m.group(0)} (sheet is {mr}r x {mc}c)")
                if tgt == name and r1 <= cell.row <= r2 and c1 <= cell.column <= c2:
                    problems["self_reference"].append(f"{name}!{cell.coordinate}: {v[:70]}")

                # blank targets. Only single cells: a range with holes in it is
                # normal — SUM and COUNT are defined over gaps — whereas a lone
                # reference to an empty cell is arithmetic on a silent zero.
                if (r1, c1) == (r2, c2) and is_empty(tgt, r1, c1):
                    if guards is None:
                        guards = guarded(bare, name)
                    if (tgt, r1, c1) not in guards:
                        problems["unguarded_blank_ref"].append(
                            f"{name}!{cell.coordinate} -> {tgt}!{get_column_letter(c1)}{r1} "
                            f"is empty: {v[:60]}")
                    else:
                        stats["blank_refs_guarded"] += 1
                stats["refs"] += 1

# ---------------------------------------------------------------- charts
RANGE = re.compile(
    r"^(?:'(?P<qs>[^']+)'|(?P<us>[^!]+))!"
    r"\$?(?P<c1>[A-Z]{1,3})\$?(?P<r1>\d+)"
    r"(?::\$?(?P<c2>[A-Z]{1,3})\$?(?P<r2>\d+))?$"
)


def chart_ref(f, tag, what):
    """Parse one chart range, reporting a bad sheet or an inverted range."""
    m = RANGE.match((f or "").strip())
    if not m:
        problems["chart_bad_ref"].append(f"{tag}: {what} is not a sheet-qualified range: {f!r}")
        return None
    tgt, r1, c1, r2, c2 = span(m, "")
    if tgt not in sheets:
        problems["chart_bad_sheet"].append(f"{tag}: {what} -> '{tgt}' — no such sheet")
        return None
    if r2 < r1 or c2 < c1:
        problems["chart_inverted_range"].append(f"{tag}: {what} -> {f} runs backwards")
        return None
    mr, mc = dims[tgt]
    if r2 > mr or c2 > mc:
        problems["chart_out_of_range"].append(
            f"{tag}: {what} -> {f} (sheet is {mr}r x {mc}c)")
    return tgt, r1, c1, r2, c2


def ref_of(part):
    """The formula string behind a chart's val / cat / tx element, if any."""
    if part is None:
        return None
    for attr in ("numRef", "strRef", "multiLvlStrRef"):
        sub = getattr(part, attr, None)
        if sub is not None and getattr(sub, "f", None):
            return sub.f
    return getattr(part, "f", None)


for name, ws in sheets.items():
    for ci, chart in enumerate(ws._charts):
        stats["charts"] += 1
        if not chart.series:
            problems["chart_no_series"].append(f"{name} chart {ci}: no series")
        for si, s in enumerate(chart.series):
            stats["chart_series"] += 1
            tag = f"{name} chart {ci} series {si}"

            val = chart_ref(ref_of(s.val), tag, "values")
            cat = chart_ref(ref_of(s.cat), tag, "categories") if s.cat is not None else None
            if s.cat is None:
                problems["chart_no_categories"].append(f"{tag}: no category range")

            if val and cat:
                nv = (val[3] - val[1] + 1) * (val[4] - val[2] + 1)
                nc = (cat[3] - cat[1] + 1) * (cat[4] - cat[2] + 1)
                if nv != nc:
                    problems["chart_length_mismatch"].append(
                        f"{tag}: {nv} values vs {nc} categories")

            # A series title is a reference to the header cell of its own column.
            # headers() returns the FIRST DATA ROW, so that header sits one row
            # above the values — an anchor off by one lands the title on a data
            # cell or on the section bar, and the legend reads as a number.
            txf = ref_of(s.tx) if s.tx is not None else None
            if txf:
                t = chart_ref(txf, tag, "title")
                if t:
                    tgt, r1, c1, r2, c2 = t
                    if (r1, c1) != (r2, c2):
                        problems["chart_title_not_single_cell"].append(f"{tag}: title -> {txf}")
                    else:
                        hv = sheets[tgt].cell(r1, c1).value
                        if not (isinstance(hv, str) and hv.strip()):
                            problems["chart_title_not_header"].append(
                                f"{tag}: title -> {tgt}!{get_column_letter(c1)}{r1} "
                                f"is {hv!r}, not a header string")
                        elif val and r1 != val[1] - 1:
                            problems["chart_title_off_header_row"].append(
                                f"{tag}: title on row {r1} but values start at row {val[1]}")

print(json.dumps({"stats": dict(stats),
                  "sheets": {k: f"{v[0]}r x {v[1]}c" for k, v in dims.items()}}, indent=1))
bad = 0
for k, vs in problems.items():
    bad += len(vs)
    print(f"\n### {k}  ({len(vs)})")
    for x in vs[:12]:
        print("   ", x)
    if len(vs) > 12:
        print(f"    ... and {len(vs)-12} more")
print("\nTOTAL PROBLEMS:", bad)
sys.exit(1 if bad else 0)
