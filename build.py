import json, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment

D = json.load(open("data.json"))
ASOF = datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")

# ---------- style kit ----------
F   = "Arial"
BLUE  = Font(name=F, size=10, color="0000FF")          # hardcoded API input
BLACK = Font(name=F, size=10)                           # formula
GREEN = Font(name=F, size=10, color="008000")           # cross-sheet link
BOLD  = Font(name=F, size=10, bold=True)
TITLE = Font(name=F, size=16, bold=True, color="1F3864")
SUB   = Font(name=F, size=9, italic=True, color="595959")
SECT  = Font(name=F, size=11, bold=True, color="FFFFFF")
HDR   = Font(name=F, size=9, bold=True, color="FFFFFF")
SECTF = PatternFill("solid", fgColor="1F3864")
HDRF  = PatternFill("solid", fgColor="4472C4")
DEMF  = PatternFill("solid", fgColor="DEEBF7")
REPF  = PatternFill("solid", fgColor="FCE4E4")
YEL   = PatternFill("solid", fgColor="FFFF00")
GREY  = PatternFill("solid", fgColor="F2F2F2")
BAND  = PatternFill("solid", fgColor="F7F9FC")
thin  = Side(style="thin", color="BFBFBF")
BOX   = Border(left=thin, right=thin, top=thin, bottom=thin)
BOT   = Border(bottom=Side(style="medium", color="1F3864"))

PCT  = '0.0%'
PCT2 = '0.00%'
DIFF = '+0.0%;-0.0%;0.0%'
USD  = '$#,##0;($#,##0);-'
NUM  = '#,##0;(#,##0);-'
DATE = 'yyyy-mm-dd'

wb = Workbook()

def title(ws, t, sub):
    ws["A1"] = t; ws["A1"].font = TITLE
    ws["A2"] = sub; ws["A2"].font = SUB

def day10(v):
    """First 10 chars of an ISO timestamp, or a placeholder when the field is null.

    These come straight from the venues. As the markets approach settlement on
    3 Nov 2026 any of them may start returning null, and None[:10] raises
    TypeError — losing the whole refresh over one cosmetic date cell.
    """
    return v[:10] if isinstance(v, str) and v else "—"


def section(ws, row, text, width):
    ws.cell(row, 1, text).font = SECT
    for c in range(1, width+1):
        ws.cell(row, c).fill = SECTF
    return row+1

def headers(ws, row, cols, widths=None):
    for i, h in enumerate(cols, 1):
        c = ws.cell(row, i, h)
        c.font = HDR; c.fill = HDRF; c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28
    return row+1

def put(ws, r, c, v, font=BLACK, fmt=None, fill=None, align=None, border=True):
    cell = ws.cell(r, c, v); cell.font = font
    if fmt: cell.number_format = fmt
    if fill: cell.fill = fill
    if align: cell.alignment = Alignment(horizontal=align)
    if border: cell.border = BOX
    return cell

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

def need(n, what):
    """Refuse to build a row range out of an empty series.

    Every LAST anchor is FIRST + len(series) - 1, so an empty series yields
    LAST < FIRST: the ref JSONs then carry an inverted range and every chart
    Reference built from the pair is silently backwards. Fail loudly instead.
    """
    if n < 1:
        raise SystemExit("FATAL: %s came back empty — refusing to write an inverted row "
                         "range. Re-run the collector before building." % what)
    return n

# ================================================================
# SHEET: Polymarket
# ================================================================
pm = wb.active; pm.title = "Polymarket"
title(pm, "Polymarket — PA-07 House Election Winner",
      "Source: Polymarket CLOB API (clob.polymarket.com) + Gamma API. Prices are per-$1 contracts; price = implied probability. Pulled %s." % ASOF)
widths(pm, {"A":13,"B":13,"C":13,"D":13,"E":13,"F":13,"G":13,"H":13,"I":13,"J":13})

r = section(pm, 4, "MARKET METADATA", 10)
r = headers(pm, r, ["Field","Democratic Party","Republican Party","","","","","","",""])
meta_rows = [
    ("Question",      D["pm_meta"]["D"]["question"],   D["pm_meta"]["R"]["question"]),
    ("Market slug",   D["pm_meta"]["D"]["slug"],       D["pm_meta"]["R"]["slug"]),
    ("CLOB token ID", D["pm_meta"]["D"]["token_id"],   D["pm_meta"]["R"]["token_id"]),
    ("Condition ID",  D["pm_meta"]["D"]["condition_id"],D["pm_meta"]["R"]["condition_id"]),
    ("Market opened", day10(D["pm_meta"]["D"]["start"]), day10(D["pm_meta"]["R"]["start"])),
    ("Resolution date", day10(D["pm_event"]["endDate"]), day10(D["pm_event"]["endDate"])),
    ("Volume ($)",    D["pm_meta"]["D"]["volume"],     D["pm_meta"]["R"]["volume"]),
    ("Liquidity ($)", D["pm_meta"]["D"]["liquidity"],  D["pm_meta"]["R"]["liquidity"]),
    ("Last trade",    D["pm_meta"]["D"]["last_trade"], D["pm_meta"]["R"]["last_trade"]),
    ("24h price change", D["pm_meta"]["D"]["one_day_change"], D["pm_meta"]["R"]["one_day_change"]),
]
PM_VOL = None
for lbl, dv, rv in meta_rows:
    put(pm, r, 1, lbl, BOLD)
    fmt = USD if "($)" in lbl else (PCT if lbl in ("Last trade",) else (DIFF if "change" in lbl else None))
    put(pm, r, 2, dv, BLUE, fmt, DEMF)
    put(pm, r, 3, rv, BLUE, fmt, REPF)
    if lbl == "Volume ($)":
        PM_VOL = r          # B = Democratic side, C = Republican side; Summary sums the pair
    r += 1
if PM_VOL is None:
    raise SystemExit("FATAL: meta_rows has no 'Volume ($)' row — the Summary sheet's "
                     "Depth / Volume cell would have nothing to point at.")

r += 1
BOOK = r
r = section(pm, r, "LIVE ORDER BOOK SNAPSHOT", 10)
r = headers(pm, r, ["Outcome","Best Bid","Best Ask","Mid","Bid-Ask Spread","Bid Size ($)","Ask Size ($)","Implied Prob (normalised)","","" ])
bk0 = r
for side, key, fill in (("D","Democratic Party",DEMF), ("R","Republican Party",REPF)):
    m = D["pm_meta"][side]
    put(pm, r, 1, key, BOLD, None, fill)
    put(pm, r, 2, m["best_bid"], BLUE, PCT)
    put(pm, r, 3, m["best_ask"], BLUE, PCT)
    put(pm, r, 4, "=(B{0}+C{0})/2".format(r), BLACK, PCT2)
    put(pm, r, 5, "=C{0}-B{0}".format(r), BLACK, DIFF)
    put(pm, r, 6, m["bid_size"], BLUE, USD)
    put(pm, r, 7, m["ask_size"], BLUE, USD)
    put(pm, r, 8, "=IFERROR(D{0}/$D${1},\"\")".format(r, bk0+2), BLACK, PCT)
    r += 1
put(pm, r, 1, "Sum of mids (100% = no vig)", BOLD, None, GREY)
put(pm, r, 4, "=SUM(D{0}:D{1})".format(bk0, bk0+1), BOLD, PCT2, YEL)
put(pm, r, 5, "=D{0}-1".format(r), BLACK, DIFF)
put(pm, r, 8, "=SUM(H{0}:H{1})".format(bk0, bk0+1), BLACK, PCT)
pm.cell(r,4).comment = Comment(
    "Two-sided overround. Sum > 100% means the pair of YES contracts costs more than the $1 "
    "they must jointly pay out; sum < 100% is a theoretical arbitrage before fees, slippage "
    "and capital cost.", "Tracker")
PM_MID_D, PM_MID_R, PM_SUM = bk0, bk0+1, r

r += 2
r = section(pm, r, "DAILY PRICE HISTORY", 10)
HH = headers(pm, r, ["Date","Dem YES Price","Rep YES Price","Sum (D+R)","Dem Normalised",
                     "Rep Normalised","Dem − Rep Spread","Dem 7d Change","Dem 30d Change",
                     "Days to Election","Axis label"])
pm.cell(HH - 1, 8).comment = Comment(
    "Change against the price N CALENDAR days earlier, not N rows earlier — the series has "
    "gaps, so the two are not the same. Where a gap leaves no quote within %d days of the "
    "target date, the cell is left blank rather than report a differently-sized window."
    % 3, "Tracker")
hist = D["pm_history"]
days = sorted(hist)
first = HH

# Tolerance, in days, on how stale a lookback baseline may be before the change
# is withheld. A 7-day change measured across a 10-day span is still a fair read;
# across 18 days it is not, and labelling it "7d" would be a lie.
CHG_TOL = 3

def change(rr, off):
    """Dem price now minus the Dem price `off` CALENDAR days earlier.

    The previous version subtracted ROWS, which is only the same thing when the
    series has no gaps — and this one has three (11, 9 and 4 days across spring
    2026). Measured against the live history that put 9 of 228 cells in the 7-day
    column and 32 of 205 in the 30-day column over the wrong window, silently:
    a plausible-looking percentage computed between the wrong two dates.

    MATCH(...,1) on the ascending date column takes the most recent row on or
    before the target, so the baseline is never newer than `off` days back. When
    a gap makes it much older than that, CHG_TOL withholds the cell rather than
    mislabel the span. #N/A from a series that does not reach back far enough
    becomes blank, which also retires the old ROW() depth guard — that one still
    wrote B0 and B-5 into the opening rows, references Excel cannot resolve.
    The `<>0` tests are the blank guard: INDEX onto an empty cell returns 0, and
    a genuine 0.000 quote does not occur here.
    """
    m = 'MATCH($A{0}-{1},$A${2}:$A{0},1)'.format(rr, off, first)
    p = 'INDEX($B${0}:$B{1},{2})'.format(first, rr, m)
    d = 'INDEX($A${0}:$A{1},{2})'.format(first, rr, m)
    return ('=IFERROR(IF(AND(ISNUMBER(B{0}),{1}<>0,{2}<>0,$A{0}-{2}<={3}),B{0}-{1},""),"")'
            .format(rr, p, d, off + CHG_TOL))
for i, day in enumerate(days):
    rr = HH + i
    dv = hist[day].get("D"); rv = hist[day].get("R")
    put(pm, rr, 1, datetime.datetime.strptime(day, "%Y-%m-%d"), BLUE, DATE)
    put(pm, rr, 2, dv, BLUE, PCT)
    put(pm, rr, 3, rv, BLUE, PCT)
    put(pm, rr, 4, "=IF(COUNT(B{0}:C{0})=2,B{0}+C{0},\"\")".format(rr), BLACK, PCT2)
    # An absent quote leaves B or C empty, and an empty cell reads as 0, not as an
    # error — so IFERROR alone lets a missing price through as 0.0%. Test the
    # numerator explicitly. (D is missing on 2026-07-29, R on 2026-04-18.)
    put(pm, rr, 5, "=IF(ISNUMBER(B{0}),IFERROR(B{0}/$D{0},\"\"),\"\")".format(rr), BLACK, PCT)
    put(pm, rr, 6, "=IF(ISNUMBER(C{0}),IFERROR(C{0}/$D{0},\"\"),\"\")".format(rr), BLACK, PCT)
    put(pm, rr, 7, "=IF(COUNT(B{0}:C{0})=2,B{0}-C{0},\"\")".format(rr), BLACK, DIFF)
    put(pm, rr, 8, change(rr, 7), BLACK, DIFF)
    put(pm, rr, 9, change(rr, 30), BLACK, DIFF)
    put(pm, rr, 10, '=DATE(2026,11,3)-A{0}'.format(rr), BLACK, NUM)
    # Short text for chart category axes. A date-valued category gets reformatted
    # and auto-sized by the renderer, which ignores the numFmt and font size set
    # on the chart itself; plain text leaves nothing to reinterpret.
    put(pm, rr, 11, datetime.datetime.strptime(day, "%Y-%m-%d").strftime("%d %b"), BLUE)
    if i % 2:
        for c in range(1, 11): pm.cell(rr, c).fill = BAND
PM_LAST = HH + need(len(days), "Polymarket daily price history") - 1
PM_FIRST = HH
pm.freeze_panes = pm.cell(HH, 2)
pm.cell(HH-1, 1).comment = Comment(
    "One row per UTC day, taken from CLOB /prices-history at fidelity=1440 (daily). "
    "Blue cells are raw API values; black cells are derived. Normalised columns rescale the "
    "pair to sum to 100%% so the two outcomes read as a clean probability split.", "Tracker")

# ================================================================
# SHEET: Kalshi
# ================================================================
ka = wb.create_sheet("Kalshi")
title(ka, "Kalshi — PA-07 House Race Winner (HOUSEPA7-26)",
      "Source: Kalshi Trade API v2 (api.elections.kalshi.com), RSA-signed requests. Prices are per-$1 contracts; price = implied probability. Pulled %s." % ASOF)
widths(ka, {"A":13,"B":13,"C":13,"D":13,"E":13,"F":13,"G":13,"H":13,"I":13,"J":13,"K":13,"L":13})

r = section(ka, 4, "MARKET METADATA", 12)
r = headers(ka, r, ["Field","Democratic (YES)","Republican (YES)","","","","","","","","",""])
km = D["k_meta"]
for lbl, dv, rv in [
    ("Market ticker", km["D"]["ticker"], km["R"]["ticker"]),
    ("Event ticker",  km["D"]["event"],  km["R"]["event"]),
    ("Question",      km["D"]["title"],  km["R"]["title"]),
    ("Candidate",     km["D"]["candidate"], km["R"]["candidate"]),
    ("Status",        km["D"]["status"], km["R"]["status"]),
    ("Opened",        day10(km["D"]["open_time"]), day10(km["R"]["open_time"])),
    ("Close time",    day10(km["D"]["close_time"]), day10(km["R"]["close_time"])),
    ("Candlesticks pulled", km["D"]["candles"], km["R"]["candles"]),
]:
    put(ka, r, 1, lbl, BOLD)
    put(ka, r, 2, dv, BLUE, NUM if "Candle" in lbl else None, DEMF)
    put(ka, r, 3, rv, BLUE, NUM if "Candle" in lbl else None, REPF)
    r += 1

r += 1
r = section(ka, r, "LIVE ORDER BOOK SNAPSHOT", 12)
r = headers(ka, r, ["Outcome","YES Best Bid","NO Best Bid","YES Best Ask (1−NO bid)","Mid",
                    "Bid-Ask Spread","YES Bid Size ($)","Implied Prob (normalised)","","","",""])
kbk = r
for side, key, fill in (("D","Democratic — Bob Brooks",DEMF), ("R","Republican — Ryan Mackenzie",REPF)):
    m = km[side]
    put(ka, r, 1, key, BOLD, None, fill)
    put(ka, r, 2, m["yes_bid"], BLUE, PCT)
    put(ka, r, 3, m["no_bid"], BLUE, PCT)
    put(ka, r, 4, "=1-C{0}".format(r), BLACK, PCT)
    put(ka, r, 5, "=(B{0}+D{0})/2".format(r), BLACK, PCT2)
    put(ka, r, 6, "=D{0}-B{0}".format(r), BLACK, DIFF)
    put(ka, r, 7, m["yes_bid_size"], BLUE, USD)
    put(ka, r, 8, "=IFERROR(E{0}/$E${1},\"\")".format(r, kbk+2), BLACK, PCT)
    r += 1
put(ka, r, 1, "Sum of mids (100% = no vig)", BOLD, None, GREY)
put(ka, r, 5, "=SUM(E{0}:E{1})".format(kbk, kbk+1), BOLD, PCT2, YEL)
put(ka, r, 6, "=E{0}-1".format(r), BLACK, DIFF)
put(ka, r, 8, "=SUM(H{0}:H{1})".format(kbk, kbk+1), BLACK, PCT)
ka.cell(kbk,4).comment = Comment(
    "Kalshi's REST market object returned null for yes_bid/yes_ask on this market, so top-of-book "
    "is derived from the /orderbook endpoint. A YES ask is the complement of the best NO bid: "
    "buying YES at price p is the same trade as selling NO at 1-p.", "Tracker")
K_MID_D, K_MID_R, K_SUM = kbk, kbk+1, r

r += 2
r = section(ka, r, "ORDER BOOK DEPTH (top 5 levels each side)", 12)
r = headers(ka, r, ["Level","Dem YES Bid","Dem YES Size ($)","Dem NO Bid","Dem NO Size ($)",
                    "Rep YES Bid","Rep YES Size ($)","Rep NO Bid","Rep NO Size ($)","","",""])
for lvl in range(5):
    put(ka, r, 1, lvl+1, BOLD, None, GREY, "center")
    for base, side in ((2,"D"), (6,"R")):
        yb = km[side]["yes_book"]; nb = km[side]["no_book"]
        put(ka, r, base,   yb[lvl][0] if lvl < len(yb) else None, BLUE, PCT)
        put(ka, r, base+1, yb[lvl][1] if lvl < len(yb) else None, BLUE, USD)
        put(ka, r, base+2, nb[lvl][0] if lvl < len(nb) else None, BLUE, PCT)
        put(ka, r, base+3, nb[lvl][1] if lvl < len(nb) else None, BLUE, USD)
    r += 1

r += 1
r = section(ka, r, "DAILY PRICE HISTORY (1-day candlesticks)", 12)
KH = headers(ka, r, ["Date","Dem YES Bid","Dem YES Ask","Dem Mid","Rep YES Bid","Rep YES Ask",
                     "Rep Mid","Sum of Mids","Dem Normalised","Dem − Rep Spread",
                     "Volume (contracts)","Open Interest","Axis label"])
kh = D["k_history"]; kdays = sorted(kh)
for i, day in enumerate(kdays):
    rr = KH + i
    dd = kh[day].get("D", {}); rr_ = kh[day].get("R", {})
    put(ka, rr, 1, datetime.datetime.strptime(day, "%Y-%m-%d"), BLUE, DATE)
    put(ka, rr, 2, dd.get("yes_bid"), BLUE, PCT)
    put(ka, rr, 3, dd.get("yes_ask"), BLUE, PCT)
    put(ka, rr, 4, "=IF(COUNT(B{0}:C{0})=2,(B{0}+C{0})/2,\"\")".format(rr), BLACK, PCT2)
    put(ka, rr, 5, rr_.get("yes_bid"), BLUE, PCT)
    put(ka, rr, 6, rr_.get("yes_ask"), BLUE, PCT)
    put(ka, rr, 7, "=IF(COUNT(E{0}:F{0})=2,(E{0}+F{0})/2,\"\")".format(rr), BLACK, PCT2)
    put(ka, rr, 8, "=IF(AND(ISNUMBER(D{0}),ISNUMBER(G{0})),D{0}+G{0},\"\")".format(rr), BLACK, PCT2)
    put(ka, rr, 9, "=IFERROR(D{0}/H{0},\"\")".format(rr), BLACK, PCT)
    put(ka, rr, 10, "=IF(AND(ISNUMBER(D{0}),ISNUMBER(G{0})),D{0}-G{0},\"\")".format(rr), BLACK, DIFF)
    put(ka, rr, 11, (dd.get("vol") or 0) + (rr_.get("vol") or 0), BLUE, NUM)
    put(ka, rr, 12, max(dd.get("oi") or 0, rr_.get("oi") or 0), BLUE, NUM)
    put(ka, rr, 13, datetime.datetime.strptime(day, "%Y-%m-%d").strftime("%d %b"), BLUE)
    if i % 2:
        for c in range(1, 14): ka.cell(rr, c).fill = BAND
K_LAST = KH + need(len(kdays), "Kalshi daily price history") - 1
K_FIRST = KH
ka.freeze_panes = ka.cell(KH, 2)

json.dump({"PM_MID_D":PM_MID_D,"PM_MID_R":PM_MID_R,"PM_SUM":PM_SUM,"PM_FIRST":PM_FIRST,"PM_LAST":PM_LAST,
           "PM_VOL":PM_VOL,
           "K_MID_D":K_MID_D,"K_MID_R":K_MID_R,"K_SUM":K_SUM,"K_FIRST":K_FIRST,"K_LAST":K_LAST},
          open("refs.json","w"))
wb.save("_stage1.xlsx")
print("stage1 ok  PM rows %d-%d  K rows %d-%d  PM volume row %d" % (PM_FIRST, PM_LAST, K_FIRST, K_LAST, PM_VOL))
