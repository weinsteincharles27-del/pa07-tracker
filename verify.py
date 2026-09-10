"""Final verification: evaluate every formula, report errors, and print each
labelled figure next to the label actually sitting in column A."""
import warnings, json
warnings.filterwarnings("ignore")
import formulas
from openpyxl import load_workbook

WB = "PA-07_House_Election_Tracker.xlsx"
xl = formulas.ExcelModel().loads(WB).finish()
sol = xl.calculate()

vals = {}
for k, v in sol.items():
    key = k.upper()
    if "'[PA-07_HOUSE_ELECTION_TRACKER.XLSX]" not in key:
        continue
    try:
        vals[key.split("]", 1)[1].replace("'", "")] = v.value[0, 0]
    except Exception:
        pass

ERRTOK = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NULL!", "#NUM!", "#CYCLE!")
errs = {}
for addr, val in vals.items():
    s = str(val)
    for t in ERRTOK:
        if t in s:
            errs.setdefault(t, []).append(addr)
total = sum(len(v) for v in errs.values())
print("cells evaluated :", len(vals))
print("ERROR CELLS     :", total)
for t, cs in errs.items():
    print("  ", t, len(cs), cs[:10])

wbf = load_workbook(WB)
print("\n=== LABELLED FIGURES (label read from column A of the same row) ===")
for sheet in ["Summary", "Margin of Victory", "PA Seat Count", "Polls", "Venue Comparison"]:
    ws = wbf[sheet]
    print("\n--- %s ---" % sheet)
    for r in range(1, min(ws.max_row, 80) + 1):
        lbl = ws.cell(r, 1).value
        f = ws.cell(r, 2).value
        if not (isinstance(lbl, str) and isinstance(f, str) and f.startswith("=")):
            continue
        v = vals.get("%s!B%d" % (sheet.upper(), r))
        if v is None:
            continue
        if isinstance(v, float):
            v = round(v, 5)
        print("   %-42s %s" % (lbl[:42], v))
