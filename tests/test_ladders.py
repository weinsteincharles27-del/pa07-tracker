"""Kalshi threshold ladders: voter turnout, and the margin ladder that is
deliberately not built.

These are survival curves, not exclusive brackets, so the arithmetic that turns
them into buckets is the whole risk surface. The margin ladder's bucket code and
its tests are in git history from before 15 Sep 2026, along with the regression
they guarded (Republican rungs differenced by row order instead of by strike).
"""
import json
import support
from openpyxl import load_workbook

WB = "PA-07_House_Election_Tracker.xlsx"


def _sheet():
    return load_workbook(support.project(WB))["Margin of Victory"]


def _k3():
    return json.load(open(support.project("k3refs.json")))


def test_the_kalshi_margin_ladder_is_collected_but_not_built():
    """Withheld on purpose: most rungs have no two-sided quote, and an expected
    margin over the one that does is a confident number from a fifth of a
    distribution. The collector keeps the history in case that changes; the
    workbook, the Summary and the charts must not show it."""
    wb = load_workbook(support.project(WB))
    mv = wb["Margin of Victory"]
    cells = " ".join(str(c.value) for row in mv.iter_rows() for c in row if c.value is not None)
    assert "KALSHI MARGIN LADDER" not in cells
    assert "Kalshi expected margin" not in cells
    sm = wb["Summary"]
    subjects = [(sm.cell(r, 1).value, sm.cell(r, 2).value) for r in range(1, sm.max_row + 1)]
    assert ("Margin of victory", "Kalshi") not in subjects
    titles = [str(ch.title.tx.rich.p[0].r[0].t) if ch.title and ch.title.tx and ch.title.tx.rich else ""
              for ws in wb.worksheets for ch in ws._charts]
    assert not any("Kalshi margin" in t for t in titles), titles
    k3 = _k3()
    for gone in ("K0", "KN", "KEM", "R03", "D03", "KHF", "KHL", "D_TOP_MID", "R_TOP_MID"):
        assert gone not in k3, gone


def test_turnout_ladder_has_a_residual_row_below_the_lowest_threshold():
    wb = load_workbook(support.project(WB))
    ws, k3 = wb["Voter Turnout"], _k3()
    f = (ws.cell(k3["T0"], 7).value or "").replace(" ", "")
    assert "1-E" in f, "the below-lowest-threshold bucket is the residual 1-P(above it): %s" % f


def test_collect3_flags_a_non_monotone_ladder():
    """P(x >= 6) can never exceed P(x >= 3). If it does, a derived bucket goes
    negative and the distribution is meaningless — say so rather than render it."""
    with support.sandbox(), support.fake_network() as up:
        up.invert_mov_rung = True
        support.run_script("collect3.py")
        d3 = support.read_json("data3.json")
    assert any(p["kind"] == "non_monotone_ladder" for p in d3["problems"]), d3["problems"]


def test_collect3_survives_a_delisted_event():
    """These settle on 3 Nov 2026; a vanished event must not kill the collector."""
    with support.sandbox(), support.fake_network() as up:
        up.missing_events = {"KXMIDTERMVOTETURN-PA07"}
        support.run_script("collect3.py")
        d3 = support.read_json("data3.json")
    assert any(p["kind"] == "missing_event" for p in d3["problems"]), d3["problems"]
    assert d3["turnout"]["rungs"] == []
    assert d3["mov_d"]["rungs"], "the other ladders should still have been collected"
